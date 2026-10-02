print("Duspul5115")

import os
import csv
from datetime import datetime

import openpyxl
import matplotlib.pyplot as plt
from openpyxl.chart import PieChart, Reference

# Change directory to the folder containing the assignment files
os.chdir(r"C:\Users\Dusti\OneDrive - ECPI University\1SDC205\Week5\5.5")


def askUser():
    total = 0

    # This loop asks the user for five numbers and adds each number
    # to the running total.
    for i in range(5):
        number = int(input("Please enter a number: "))
        total += number

    print("The sum for the 5 numbers entered is:", total)


def askIncome():
    # This loop asks for the name and annual income of five people
    # and appends each entry to the existing final.csv file.
    for i in range(5):
        name = input("Please enter a name: ")
        income = int(input("Please enter their income: "))

        with open("final.csv", "a", newline="") as file:
            writer = csv.writer(file)
            writer.writerow([name, income])


def excelPie():
    names = []
    incomes = []

    # Read all names and incomes from final.csv
    with open("final.csv", "r") as file:
        reader = csv.reader(file)

        for row in reader:
            if row:
                names.append(row[0])
                incomes.append(int(row[1]))

    # Create a new Excel workbook
    wb = openpyxl.Workbook()

    # Select the active worksheet
    ws = wb.active

    # Add each name and income to the Excel worksheet
    for i in range(len(names)):
        ws.cell(row=i + 1, column=1, value=names[i])
        ws.cell(row=i + 1, column=2, value=incomes[i])

    # Create a new pie chart
    pieChart = PieChart()

    # Select the numerical income values for the pie chart
    data = Reference(ws, min_col=2, min_row=1, max_row=len(incomes))

    # Select the names that will label each section of the pie chart
    labels = Reference(ws, min_col=1, min_row=1, max_row=len(names))

    # Add the income data to the pie chart
    pieChart.add_data(data)

    # Add the names as the pie chart categories
    pieChart.set_categories(labels)

    # Create today's date for the chart title
    today = datetime.now().strftime("%B %d, %Y")

    # Set the pie chart title to the Student ID and today's date
    pieChart.title = "Duspul5115 " + today

    # Set the pie chart size
    pieChart.height = 10
    pieChart.width = 15

    # Add the pie chart to the worksheet
    ws.add_chart(pieChart, "A11")

    # Save the worksheet and pie chart as final.xlsx
    wb.save("final.xlsx")


def verticalBar():
    names = []
    incomes = []

    # Read all names and incomes from final.csv
    with open("final.csv", "r") as file:
        reader = csv.reader(file)

        for row in reader:
            if row:
                names.append(row[0])
                incomes.append(int(row[1]))

    # Create today's date for the chart title
    today = datetime.now().strftime("%B %d, %Y")

    # Create a vertical bar graph using the names and incomes
    plt.bar(names, incomes, label="Income")

    # Label the horizontal axis
    plt.xlabel("Name")

    # Label the vertical axis
    plt.ylabel("Income")

    # Set the graph title to the Student ID and today's date
    plt.title("Duspul5115 " + today)

    # Display the graph legend
    plt.legend()

    # Display the completed vertical bar graph
    plt.show()


askUser()
askIncome()
excelPie()
verticalBar()
